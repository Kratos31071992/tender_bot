import asyncio
import logging
import os
import re
import zipfile
from pathlib import Path

from aiogram import Bot, Dispatcher, F, types
from aiogram.filters import Command, CommandStart
from aiogram.types import FSInputFile, InlineKeyboardButton, InlineKeyboardMarkup
from doc2txt import extract_text
from docx import Document
from dotenv import load_dotenv
from openpyxl import load_workbook
from pypdf import PdfReader

from analyze import AnalyzeError, analyze_tender
from database import (
    add_demo_tender,
    get_history,
    init_db,
    list_tenders,
    log_search,
    save_and_should_show,
    save_tender,
)
from parser import (
    download_bidzaar_files,
    parse_bidzaar,
    safe_filename,
    search_bidzaar,
)


load_dotenv()
token = os.getenv('BOT_TOKEN')
bot = Bot(token=token)
dp = Dispatcher()

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger('tender_bot')

DOCS_DIR = Path(__file__).resolve().parent / 'docs'
LAST_TENDER = {}
WATCH = {}
REFRESH_SECONDS = 15 * 60
CARD_PAUSE = 1.5


def read_tender_file(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == '.docx':
            word = Document(path)
            return '\n'.join(p.text for p in word.paragraphs if p.text.strip())
        if suffix == '.doc':
            return extract_text(str(path)) or ''
        if suffix == '.pdf':
            reader = PdfReader(str(path))
            pages = []
            for page in reader.pages:
                chunk = page.extract_text() or ''
                if chunk.strip():
                    pages.append(chunk)
            return '\n'.join(pages)
        if suffix == '.xlsx':
            book = load_workbook(str(path), data_only=True)
            lines = []
            for sheet in book.worksheets:
                lines.append(f'лист: {sheet.title}')
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(cell) for cell in row if cell is not None]
                    if cells:
                        lines.append('\t'.join(cells))
            return '\n'.join(lines)
        if suffix == '.zip':
            out_dir = path.parent / (path.stem + '_unzip')
            out_dir.mkdir(exist_ok=True)
            with zipfile.ZipFile(path) as archive:
                archive.extractall(out_dir)
            parts = []
            for inner in out_dir.rglob('*'):
                if not inner.is_file() or inner.suffix.lower() == '.zip':
                    continue
                text = read_tender_file(inner)
                if text.strip():
                    parts.append(f'--- {inner.name} ---\n{text}')
            return '\n'.join(parts)
        return ''
    except Exception as error:
        return f'[не прочитал {path.name}: {error}]'


def folder_name(data: dict) -> str:
    title = data.get('title') or 'tender'
    short = re.sub(r'[^\wа-яА-ЯёЁ\s-]', '', title)
    short = '_'.join(short.split())[:40] or 'tender'
    category = data.get('category') or 'разное'
    return f"{category}__{short}__{data['id'][:8]}"


def card_keyboard(tender_id: str) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[[
            InlineKeyboardButton(
                text='Скачать файлы',
                callback_data=f'dl:{tender_id}',
            ),
            InlineKeyboardButton(
                text='Анализ',
                callback_data=f'an:{tender_id}',
            ),
        ]]
    )


def remember_folder(user_id: int, tender_id: str, folder: str) -> None:
    LAST_TENDER[user_id] = folder
    LAST_TENDER[f'{user_id}:{tender_id}'] = folder


def folder_for(user_id: int, tender_id=None):
    if tender_id:
        return LAST_TENDER.get(f'{user_id}:{tender_id}') or LAST_TENDER.get(user_id)
    return LAST_TENDER.get(user_id)


async def send_card(chat_id: int, row: dict) -> None:
    await bot.send_message(
        chat_id,
        f'{row["title"]}\n{row["url"]}',
        reply_markup=card_keyboard(row['id']),
    )


async def send_new_cards(chat_id: int, user_id: int, rows: list) -> int:
    shown = 0
    for row in rows:
        if not await asyncio.to_thread(save_and_should_show, user_id, row):
            continue
        await send_card(chat_id, row)
        shown += 1
        await asyncio.sleep(CARD_PAUSE)
    return shown


def fetch_and_save_files(url: str, user_id: int):
    data = parse_bidzaar(url)
    folder = folder_name(data)
    folder_path = DOCS_DIR / folder
    saved = []
    if data['files']:
        saved = download_bidzaar_files(data['files'], str(folder_path))
    remember_folder(user_id, data['id'], folder)
    save_tender(data)
    return data, saved, folder_path


def collect_folder_text(folder_path: Path) -> tuple[str, list[str]]:
    pieces = []
    skipped = []
    for path in folder_path.iterdir():
        if not path.is_file() or path.name == 'analiz.txt':
            continue
        text = read_tender_file(path)
        if text.strip() and not text.startswith('[не прочитал'):
            pieces.append(f'--- {path.name} ---\n{text}')
        else:
            skipped.append(path.name)
    return '\n'.join(pieces), skipped


async def run_analyze(message_or_cb, folder_path: Path) -> None:
    send = message_or_cb.answer
    if not folder_path.is_dir():
        await send(f'Папки нет: {folder_path.name}')
        return

    body, skipped = await asyncio.to_thread(collect_folder_text, folder_path)
    if not body:
        await send(
            'В папке нет текста для анализа. '
            f'Пропущены: {", ".join(skipped) or "ничего"}'
        )
        return

    await send('Читаю документы, жду Яндекс…')
    try:
        report = await asyncio.to_thread(analyze_tender, body)
    except AnalyzeError as error:
        logger.warning('Анализ не удался: %s', error)
        await send(str(error))
        return

    analiz_path = folder_path / 'analiz.txt'
    analiz_path.write_text(report or 'Пустой анализ', encoding='utf-8')
    await message_or_cb.answer_document(
        FSInputFile(analiz_path),
        caption='Анализ закупки',
    )


def stop_watch(user_id: int) -> None:
    job = WATCH.pop(user_id, None)
    if job and job.get('task'):
        job['task'].cancel()


async def watch_loop(user_id: int, chat_id: int, query: str) -> None:
    try:
        while True:
            await asyncio.sleep(REFRESH_SECONDS)
            try:
                rows = await asyncio.to_thread(search_bidzaar, query)
                shown = await send_new_cards(chat_id, user_id, rows)
            except ConnectionError:
                try:
                    await bot.send_message(
                        chat_id, 'Обновление: Bidzaar не ответил.'
                    )
                except Exception:
                    logger.exception('watch_loop: не смог написать про Bidzaar')
                continue
            except Exception:
                logger.exception(
                    'watch_loop: сбой обновления по «%s»', query
                )
                continue
            if shown:
                try:
                    await bot.send_message(
                        chat_id,
                        f'Обновление по «{query}»: новых {shown}.',
                    )
                except Exception:
                    logger.exception('watch_loop: не смог отправить счётчик')
    except asyncio.CancelledError:
        return


def start_watch(user_id: int, chat_id: int, query: str) -> None:
    stop_watch(user_id)
    WATCH[user_id] = {
        'query': query,
        'task': asyncio.create_task(watch_loop(user_id, chat_id, query)),
    }


@dp.message(CommandStart())
async def start(message: types.Message):
    await message.answer(
        'Бот мониторинга закупок запущен.\n'
        '/fresh слово — поиск на Bidzaar\n'
        '/parse ссылка — карточка и файлы\n'
        '/analyze — анализ последней папки\n'
        '/history — история твоих поисков\n'
        '/stop — остановить обновление'
    )


@dp.message(Command('demo'))
async def demo(message: types.Message):
    await asyncio.to_thread(add_demo_tender)
    rows = await asyncio.to_thread(list_tenders)
    lines = [f'{row[0]} | {row[1]} | {row[2]}' for row in rows]
    await message.answer('\n'.join(lines) or 'База пустая')


@dp.message(Command('parse'))
async def cmd_parse(message: types.Message):
    parts = message.text.split(maxsplit=1)
    if len(parts) < 2:
        await message.answer('Кинь ссылку: /parse https://bidzaar.com/...')
        return
    url = parts[1]
    await asyncio.to_thread(log_search, message.from_user.id, '/parse', url)
    await message.answer('Читаю карточку Bidzaar…')
    try:
        data, saved, folder_path = await asyncio.to_thread(
            fetch_and_save_files,
            url,
            message.from_user.id,
        )
    except (ValueError, ConnectionError) as error:
        await message.answer(str(error))
        return

    lines = [data['title'], data['url'], '', 'Документы:']
    if data['files']:
        for i, file in enumerate(data['files'], start=1):
            name = file.get('name') or 'без имени'
            ext = file.get('extension') or ''
            lines.append(f'{i}. {name}.{ext}' if ext else f'{i}. {name}')
        lines.append('')
        lines.append(f'Скачал {len(saved)} файл(ов) в {folder_path}')
    else:
        lines.append('в карточке файлов нет')
    await message.answer('\n'.join(lines))


@dp.message(Command('fresh'))
async def fresh_cmd(message: types.Message):
    query = message.text.replace('/fresh', '', 1).strip()
    if not query:
        await message.answer('Напиши слово после команды.\nПример: /fresh клининг')
        return
    await asyncio.to_thread(log_search, message.from_user.id, '/fresh', query)
    await message.answer(f'Ищу на Bidzaar: {query}')
    try:
        rows = await asyncio.to_thread(search_bidzaar, query)
    except ConnectionError:
        await message.answer('Bidzaar не ответил. Попробуй ещё раз.')
        return
    if not rows:
        await message.answer('Ничего не нашёл по этому запросу.')
        return
    shown = await send_new_cards(message.chat.id, message.from_user.id, rows)
    if shown == 0:
        await message.answer(
            'Тебе эти закупки уже показывали. Новых карточек нет.'
        )
    start_watch(message.from_user.id, message.chat.id, query)
    await message.answer(
        f'Смотрю «{query}» каждые 15 минут. Новое слово — новый /fresh. Стоп — /stop'
    )


@dp.message(Command('stop'))
async def cmd_stop(message: types.Message):
    stop_watch(message.from_user.id)
    await message.answer('Обновление остановлено.')


@dp.message(Command('history'))
async def cmd_history(message: types.Message):
    rows = await asyncio.to_thread(get_history, message.from_user.id)
    if not rows:
        await message.answer('Истории пока нет. Сделай /fresh или /parse.')
        return
    lines = ['Последние запросы:']
    for row in rows:
        lines.append(f'{row["created_at"]}  {row["command"]}  {row["query"]}')
    await message.answer('\n'.join(lines))


@dp.message(Command('analyze'))
async def cmd_analyze(message: types.Message):
    folder = folder_for(message.from_user.id)
    if not folder:
        await message.answer('Сначала скачай файлы с карточки или /parse')
        return
    await run_analyze(message, DOCS_DIR / folder)


@dp.callback_query(F.data.startswith('dl:'))
async def cb_download(callback: types.CallbackQuery):
    tender_id = callback.data.split(':', 1)[1]
    url = f'https://bidzaar.com/app/process/light/{tender_id}'
    try:
        await callback.answer('Качаю файлы…')
    except Exception:
        pass
    try:
        data, saved, folder_path = await asyncio.to_thread(
            fetch_and_save_files,
            url,
            callback.from_user.id,
        )
    except (ValueError, ConnectionError) as error:
        await callback.message.answer(str(error))
        return
    if not data['files']:
        await callback.message.answer('В карточке файлов нет')
        return
    await callback.message.answer(
        f'Скачал {len(saved)} файл(ов) в {folder_path}'
    )


@dp.callback_query(F.data.startswith('an:'))
async def cb_analyze(callback: types.CallbackQuery):
    tender_id = callback.data.split(':', 1)[1]
    url = f'https://bidzaar.com/app/process/light/{tender_id}'
    try:
        await callback.answer('Готовлю анализ…')
    except Exception:
        pass
    folder = folder_for(callback.from_user.id, tender_id)
    folder_path = DOCS_DIR / folder if folder else None
    if folder_path is None or not folder_path.is_dir():
        try:
            _data, _saved, folder_path = await asyncio.to_thread(
                fetch_and_save_files,
                url,
                callback.from_user.id,
            )
        except (ValueError, ConnectionError) as error:
            await callback.message.answer(str(error))
            return
    await run_analyze(callback.message, folder_path)


@dp.message(F.document)
async def save_doc(message: types.Message):
    doc = message.document
    if not doc.file_name:
        await message.answer('У файла нет имени, не сохраняю')
        return
    user_id = message.from_user.id
    tender_folder = folder_for(user_id) or 'без_закупки'
    user_dir = DOCS_DIR / str(user_id) / tender_folder
    user_dir.mkdir(parents=True, exist_ok=True)
    dest = user_dir / safe_filename(doc.file_name)
    await bot.download(doc, destination=dest)

    text = await asyncio.to_thread(read_tender_file, dest)
    reply = [f'Принял: {dest.name}', f'Папка: {user_dir}']
    if not text.strip() or text.startswith('[не прочитал'):
        reply.append('Текст не вытащил (пустой файл или формат пока не читаем)')
        await message.answer('\n'.join(reply))
        return

    await message.answer('Читаю документ, жду Яндекс…')
    try:
        report = await asyncio.to_thread(analyze_tender, text)
    except AnalyzeError as error:
        logger.warning('Анализ файла не удался: %s', error)
        await message.answer(str(error))
        return

    analiz_path = user_dir / 'analiz.txt'
    analiz_path.write_text(report or 'Пустой анализ', encoding='utf-8')
    await message.answer_document(
        FSInputFile(analiz_path),
        caption='Анализ закупки',
    )
    await message.answer('\n'.join(reply))


async def main():
    init_db()
    await dp.start_polling(bot)


if __name__ == '__main__':
    asyncio.run(main())
